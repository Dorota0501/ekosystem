from __future__ import annotations

import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

if __package__:
    from .KNNImputer import KNNDataImputer
else:
    from KNNImputer import KNNDataImputer


def _unique_path(path: Path) -> Path:
    """Return a non-existing path by adding a numeric suffix when necessary.

    The original path is returned when it does not exist. Otherwise ``_1``,
    ``_2`` and subsequent numbers are appended before the file extension.

    :param path: Preferred output path.
    :type path: pathlib.Path
    :return: Available output path.
    :rtype: pathlib.Path
    """
    if not path.exists():
        return path

    number = 1
    while True:
        candidate = path.with_name(f"{path.stem}_{number}{path.suffix}")
        if not candidate.exists():
            return candidate
        number += 1


def _highlight_missing_cells(
        workbook_path: Path,
        missing_records: list[dict[str, object]],
        scope: str | None = None,
) -> None:
    """Highlight cells that were missing before KNN imputation.

    The same yellow fill is used in the simulated workbook and in completed
    KNN output workbooks. In a simulated workbook the highlighted cell remains
    empty. In a completed workbook the cell contains the imputed value but
    keeps the yellow fill, making every imputed position easy to identify.

    For ``ALL`` output and simulated input files, original row positions are
    preserved and can be used directly. For ``LS`` and ``NL`` outputs, rows
    are filtered after imputation, so objects are located by their ``Name``
    value instead.

    :param workbook_path: Workbook whose missing or imputed cells should be
        highlighted.
    :type workbook_path: pathlib.Path
    :param missing_records: Detailed records describing original and simulated
        missing cells.
    :type missing_records: list[dict[str, object]]
    :param scope: Optional output scope. Use ``None`` for a simulated input
        workbook or ``ALL``, ``LS`` or ``NL`` for an imputed workbook.
    :type scope: str or None
    :return: None
    :rtype: None
    """
    workbook = load_workbook(workbook_path)
    yellow_fill = PatternFill(fill_type="solid", fgColor="FFFF00")

    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        sheet_records = [
            record
            for record in missing_records
            if str(record["Sheet_Name"]) == sheet_name
            and (
                scope in {None, "ALL"}
                or str(record.get("LS_NL", "")).upper() == scope
            )
        ]

        if not sheet_records:
            continue

        # Mapa nazw kolumn pozwala znaleźć atrybut niezależnie od jego numeru.
        header_map = {
            str(cell.value).strip(): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }

        name_column = header_map.get("Name")
        name_to_rows: dict[str, list[int]] = {}

        # Po podziale LS/NL indeksy wierszy zmieniają się, dlatego obiekty
        # w tych plikach są wyszukiwane na podstawie kolumny Name.
        if scope in {"LS", "NL"} and name_column is not None:
            for row_number in range(2, worksheet.max_row + 1):
                value = worksheet.cell(row=row_number, column=name_column).value
                if value is None:
                    continue
                name_to_rows.setdefault(str(value).strip(), []).append(row_number)

        for record in sheet_records:
            attribute = str(record["Attribute"])
            column_number = header_map.get(attribute)
            if column_number is None:
                continue

            if scope in {"LS", "NL"}:
                if name_column is None:
                    continue
                row_numbers = name_to_rows.get(str(record["Name"]).strip(), [])
            else:
                # W pliku z symulowanymi brakami i w ALL zachowana jest
                # oryginalna kolejność wierszy; +2 uwzględnia nagłówek Excela.
                row_numbers = [int(record["Row_Index"]) + 2]

            for row_number in row_numbers:
                if 2 <= row_number <= worksheet.max_row:
                    worksheet.cell(
                        row=row_number,
                        column=column_number,
                    ).fill = yellow_fill

    workbook.save(workbook_path)


def main() -> None:
    """Run missing-data simulation and KNN imputation for project datasets.

    For every configured missing-data percentage, the procedure creates an
    artificial test workbook by hiding one observed ecological feature in the
    requested percentage of ``LS`` objects and independently in the requested
    percentage of ``NL`` objects from every biological dataset. Original
    missing values are preserved and reported separately from simulated ones.

    The damaged workbook is then processed by :class:`KNNDataImputer`. KNN
    imputation is performed on all objects from each biological dataset before
    the completed data are split into ``ALL``, ``LS`` and ``NL`` outputs.

    All files generated during one execution are stored in a dedicated run
    directory. Simulated input workbooks use names such as
    ``missing_5_dane.xlsx`` and completed workbooks use names such as
    ``Fill_missing_5_dane_ALL.xlsx``. When a preferred file name already
    exists, a numeric suffix is added before the extension. Missing cells are
    highlighted in yellow in simulated workbooks. The same cells remain yellow
    after KNN fills them in completed ALL/LS/NL workbooks.

    The shared ``KNN_parameters.xlsx`` file remains directly in the main KNN
    results directory and is updated by :class:`KNNDataImputer`. Detailed
    information about original and simulated missing cells is written to
    ``KNN_imputation_details.xlsx`` in the run directory.

    For simulated missing values, the original value is known. Therefore the
    detailed report also stores the absolute error between the original and
    imputed values. No error can be calculated for values that were already
    missing in the source workbook.

    Optionally, the generated ``ALL`` workbook can temporarily replace the
    project ``dane.xlsx`` file while additional decision scripts are executed.
    The original workbook is restored afterwards, including when one of the
    additional scripts fails.

    :return: None
    :rtype: None
    :raises FileNotFoundError: If the source workbook, LS/NL membership file or
        an enabled additional script cannot be found.
    :raises ValueError: If a configured missing-data percentage is invalid or
        a dataset contains too few eligible objects for the simulation.
    :raises RuntimeError: If the ALL output cannot be identified uniquely or an
        enabled additional script finishes with a non-zero exit code.
    """
    current_dir = Path(__file__).resolve().parent
    project_root = current_dir.parent

    input_file = project_root / "dane.xlsx"
    forest_file = project_root / "Udzial lasu.xlsx"
    output_directory = project_root / "results_knn_imputer"
    output_directory.mkdir(parents=True, exist_ok=True)

    # Jeden katalog przechowuje wszystkie pliki wygenerowane podczas jednego
    # uruchomienia. Wspólny KNN_parameters.xlsx pozostaje katalog poziom wyżej.
    run_datetime = datetime.now()
    run_signature = run_datetime.strftime("%Y%m%d_%H%M%S_%f")
    run_directory = output_directory / f"ALL_{run_signature}"
    run_directory.mkdir(parents=True, exist_ok=False)

    # Parametry algorytmu KNN.
    k_neighbour = 5
    metric = "euclidean"           # "euclidean", "manhattan" lub własna funkcja
    metric_direction = "min"       # "min" dla odległości, "max" dla podobieństwa

    # Parametry symulowania braków danych.
    missing_percentages = (0.05, 0.25, 0.50)
    allow_zero_values = False
    random_state = 42

    # Zakresy danych zapisywane po wykonaniu jednej imputacji pełnego zbioru.
    scopes = ("ALL", "LS", "NL")

    # Opcjonalne uruchomienie istniejących skryptów decyzyjnych po imputacji.
    # Włączenie wymaga ustawienia run_decision_scripts = True i podania ścieżek.
    run_decision_scripts = False
    decision_scripts = (
        # project_root / "main.py",
        # project_root / "odczyt.py",
        # project_root / "read_data.py",
    )

    if not input_file.exists():
        raise FileNotFoundError(f"Input workbook not found: {input_file}")

    if any(scope in {"LS", "NL"} for scope in scopes) and not forest_file.exists():
        raise FileNotFoundError(
            f"LS/NL membership file not found: {forest_file}. "
            "It is required when LS or NL output is requested."
        )

    # Wczytanie przypisania obiektów do LS/NL pozwala zapisać tę informację
    # również przy każdym szczegółowo raportowanym braku danych.
    membership = KNNDataImputer._load_membership(forest_file)
    membership_map = dict(
        zip(
            membership["Name"].astype(str),
            membership["LS/NL"].astype(str),
        )
    )

    imputer = KNNDataImputer(
        n_neighbors=k_neighbour,
        metric=metric,
        metric_direction=metric_direction,
        output_directory=run_directory,
        parameter_directory=output_directory,
    )

    details_path = run_directory / "KNN_imputation_details.xlsx"
    details_columns = [
        "Run_DateTime",
        "Missing_Percentage",
        "Dataset",
        "Sheet_Name",
        "Name",
        "LS_NL",
        "Row_Index",
        "Attribute",
        "Missing_Source",
        "Original_Value",
        "Imputed_Value",
        "Absolute_Error",
        "Algorithm",
        "Metric",
        "Metric_Direction",
        "k_Neighbour",
        "Allow_Zero_Values",
        "Random_State",
        "Simulated_Input_File",
        "Imputed_Output_File",
    ]

    print("\n========== KNN RUN START ==========")
    print(f"Input file: {input_file}")
    print(f"Results root: {output_directory}")
    print(f"Run directory: {run_directory}")
    print(f"k: {k_neighbour}")
    print(f"Metric: {getattr(metric, '__name__', metric)} ({metric_direction})")
    print(f"Missing percentages: {missing_percentages}")

    for percentage_index, missing_percentage in enumerate(missing_percentages):
        if not 0 < missing_percentage <= 1:
            raise ValueError(
                "Every missing-data percentage must be in the interval (0, 1]."
            )

        simulation_datetime = datetime.now()
        simulation_signature = simulation_datetime.strftime("%Y%m%d_%H%M%S_%f")
        percentage_label = f"{missing_percentage * 100:g}".replace(".", "p")
        simulation_random_state = random_state + percentage_index

        print(f"\n----- Simulation {missing_percentage * 100:.1f}% -----")

        rng = np.random.default_rng(simulation_random_state)

        excel = pd.ExcelFile(input_file, engine="openpyxl")
        damaged_sheets: dict[str, pd.DataFrame] = {}
        missing_records: list[dict[str, object]] = []

        for sheet_name in excel.sheet_names:
            df = pd.read_excel(input_file, sheet_name=sheet_name, engine="openpyxl")
            features = imputer.feature_columns(df)

            # Arkusze bez atrybutów ekologicznych są kopiowane bez zmian.
            if not features:
                damaged_sheets[sheet_name] = df.copy()
                continue

            dataset = imputer._dataset_name(sheet_name)
            damaged = df.copy(deep=True)
            numeric = damaged[features].apply(pd.to_numeric, errors="coerce")

            # Zapisanie braków obecnych już w oryginalnym pliku wejściowym.
            for row_pos in range(len(damaged)):
                object_name = (
                    damaged.iloc[row_pos]["Name"]
                    if "Name" in damaged.columns
                    else row_pos
                )
                object_scope = membership_map.get(str(object_name), "")

                for feature in features:
                    if pd.isna(numeric.iloc[row_pos][feature]):
                        missing_records.append(
                            {
                                "Run_DateTime": simulation_datetime.strftime(
                                    "%Y-%m-%d %H:%M:%S.%f"
                                ),
                                "Missing_Percentage": missing_percentage,
                                "Dataset": dataset,
                                "Sheet_Name": sheet_name,
                                "Name": object_name,
                                "LS_NL": object_scope,
                                "Row_Index": row_pos,
                                "Attribute": feature,
                                "Missing_Source": "ORIGINAL",
                                "Original_Value": np.nan,
                            }
                        )

            # Obiekty kwalifikujące się do symulacji są rozdzielane na LS i NL.
            # Dzięki temu zadany procent braków jest losowany niezależnie w
            # każdej z tych dwóch grup, a nie z całego arkusza jednocześnie.
            eligible_rows_by_scope: dict[str, list[tuple[int, list[str]]]] = {
                "LS": [],
                "NL": [],
            }
            total_objects_by_scope = {"LS": 0, "NL": 0}

            for row_pos in range(len(damaged)):
                object_name = (
                    damaged.iloc[row_pos]["Name"]
                    if "Name" in damaged.columns
                    else row_pos
                )
                object_scope = membership_map.get(str(object_name), "").upper()

                if object_scope not in {"LS", "NL"}:
                    continue

                total_objects_by_scope[object_scope] += 1
                eligible_features: list[str] = []

                for feature in features:
                    value = numeric.iloc[row_pos][feature]
                    if pd.isna(value):
                        continue
                    if not allow_zero_values and float(value) == 0.0:
                        continue
                    eligible_features.append(feature)

                if eligible_features:
                    eligible_rows_by_scope[object_scope].append(
                        (row_pos, eligible_features)
                    )

            simulated_counts: dict[str, int] = {"LS": 0, "NL": 0}

            for object_scope in ("LS", "NL"):
                group_size = total_objects_by_scope[object_scope]
                eligible_rows = eligible_rows_by_scope[object_scope]

                if group_size == 0:
                    print(f"[{dataset}] {object_scope}: no objects found")
                    continue

                # Procent liczony jest względem wszystkich obiektów należących
                # do danej grupy LS/NL, a nie względem całego arkusza.
                objects_to_damage = max(
                    1,
                    int(round(group_size * missing_percentage)),
                )

                if objects_to_damage > len(eligible_rows):
                    raise ValueError(
                        f"Dataset '{dataset}', scope '{object_scope}' has only "
                        f"{len(eligible_rows)} eligible objects, but "
                        f"{objects_to_damage} were requested for the "
                        f"{missing_percentage:.2%} simulation of "
                        f"{group_size} {object_scope} objects."
                    )

                selected_positions = rng.choice(
                    len(eligible_rows),
                    size=objects_to_damage,
                    replace=False,
                )

                for selected_position in selected_positions:
                    row_pos, eligible_features = eligible_rows[
                        int(selected_position)
                    ]
                    feature = str(rng.choice(eligible_features))
                    original_value = float(
                        pd.to_numeric(damaged.iloc[row_pos][feature])
                    )
                    object_name = (
                        damaged.iloc[row_pos]["Name"]
                        if "Name" in damaged.columns
                        else row_pos
                    )

                    missing_records.append(
                        {
                            "Run_DateTime": simulation_datetime.strftime(
                                "%Y-%m-%d %H:%M:%S.%f"
                            ),
                            "Missing_Percentage": missing_percentage,
                            "Dataset": dataset,
                            "Sheet_Name": sheet_name,
                            "Name": object_name,
                            "LS_NL": object_scope,
                            "Row_Index": row_pos,
                            "Attribute": feature,
                            "Missing_Source": "SIMULATED",
                            "Original_Value": original_value,
                        }
                    )

                    damaged.at[damaged.index[row_pos], feature] = np.nan

                simulated_counts[object_scope] = objects_to_damage

            damaged_sheets[sheet_name] = damaged
            total_simulated = simulated_counts["LS"] + simulated_counts["NL"]
            print(
                f"[{dataset}] simulated missing values: "
                f"LS={simulated_counts['LS']} of {total_objects_by_scope['LS']}, "
                f"NL={simulated_counts['NL']} of {total_objects_by_scope['NL']}, "
                f"ALL={total_simulated} of {len(damaged)}"
            )

        simulated_input_path = _unique_path(
            run_directory / f"missing_{percentage_label}_{input_file.name}"
        )

        # Zapisanie wersji danych zawierającej braki oryginalne oraz symulowane.
        with pd.ExcelWriter(simulated_input_path, engine="openpyxl") as writer:
            for sheet_name in excel.sheet_names:
                damaged_sheets[sheet_name].to_excel(
                    writer,
                    sheet_name=sheet_name,
                    index=False,
                )

        # Wszystkie komórki zawierające braki, zarówno oryginalne jak i
        # symulowane, są zaznaczane na żółto w pliku z brakami.
        _highlight_missing_cells(
            workbook_path=simulated_input_path,
            missing_records=missing_records,
        )

        print(f"Saved simulated missing-data workbook: {simulated_input_path}")

        # Imputacja pełnych grup biologicznych, zapis ALL/LS/NL oraz aktualizacja
        # wspólnego pliku KNN_parameters.xlsx.
        simulated_sample_name = simulated_input_path.stem
        if simulated_sample_name.startswith("missing_"):
            simulated_sample_name = simulated_sample_name[len("missing_"):]

        summaries = imputer.impute_workbook(
            input_file=simulated_input_path,
            forest_file=forest_file,
            scopes=scopes,
            output_base_name=f"Fill_missing_{simulated_sample_name}",
        )

        output_paths_by_scope = {
            summary.scope: run_directory / summary.output_file
            for summary in summaries
        }

        # Te same komórki pozostają zaznaczone na żółto po imputacji.
        # Wartość jest już uzupełniona przez KNN, ale kolor pozwala łatwo
        # odróżnić ją od wartości obecnych w danych wejściowych.
        for scope in scopes:
            output_path = output_paths_by_scope.get(scope)
            if output_path is not None:
                _highlight_missing_cells(
                    workbook_path=output_path,
                    missing_records=missing_records,
                    scope=scope,
                )

        all_output_names = {
            summary.output_file
            for summary in summaries
            if summary.scope == "ALL"
        }
        if len(all_output_names) != 1:
            raise RuntimeError(
                "Could not uniquely determine the ALL imputed output workbook."
            )

        all_output_path = run_directory / next(iter(all_output_names))
        imputed_excel = pd.ExcelFile(all_output_path, engine="openpyxl")
        imputed_sheets = {
            sheet_name: pd.read_excel(
                all_output_path,
                sheet_name=sheet_name,
                engine="openpyxl",
            )
            for sheet_name in imputed_excel.sheet_names
        }

        # Uzupełnienie raportu o wartości wyznaczone przez KNN i, w przypadku
        # braków symulowanych, obliczenie bezwzględnego błędu imputacji.
        completed_records: list[dict[str, object]] = []
        metric_name = imputer._metric_name()

        for record in missing_records:
            sheet_name = str(record["Sheet_Name"])
            feature = str(record["Attribute"])
            row_pos = int(record["Row_Index"])
            imputed_value = float(
                pd.to_numeric(imputed_sheets[sheet_name].iloc[row_pos][feature])
            )

            original_value = record["Original_Value"]
            absolute_error = (
                abs(imputed_value - float(original_value))
                if record["Missing_Source"] == "SIMULATED"
                else np.nan
            )

            completed_records.append(
                {
                    **record,
                    "Imputed_Value": imputed_value,
                    "Absolute_Error": absolute_error,
                    "Algorithm": "KNN Imputer",
                    "Metric": metric_name,
                    "Metric_Direction": metric_direction,
                    "k_Neighbour": k_neighbour,
                    "Allow_Zero_Values": allow_zero_values,
                    "Random_State": simulation_random_state,
                    "Simulated_Input_File": str(simulated_input_path.resolve()),
                    "Imputed_Output_File": str(all_output_path.resolve()),
                }
            )

        new_details = pd.DataFrame(completed_records, columns=details_columns)

        # Wszystkie poziomy braków z jednego uruchomienia są dopisywane do
        # jednego raportu szczegółowego w katalogu bieżącego uruchomienia.
        if details_path.exists():
            previous_details = pd.read_excel(
                details_path,
                sheet_name="Imputation_Details",
                engine="openpyxl",
            ).reindex(columns=details_columns)
            details_to_save = pd.concat(
                [previous_details, new_details],
                ignore_index=True,
            )
        else:
            details_to_save = new_details

        with pd.ExcelWriter(details_path, engine="openpyxl") as writer:
            details_to_save.to_excel(
                writer,
                sheet_name="Imputation_Details",
                index=False,
            )

        print(f"Updated imputation details: {details_path}")
        print(f"ALL workbook for decision code: {all_output_path}")

        # Opcjonalne uruchomienie istniejących skryptów decyzyjnych. Na czas
        # obliczeń uzupełniony plik ALL zastępuje dane.xlsx, po czym oryginalny
        # plik jest przywracany niezależnie od wyniku dodatkowych skryptów.
        if run_decision_scripts:
            if not decision_scripts:
                raise RuntimeError(
                    "run_decision_scripts=True, but decision_scripts is empty."
                )

            missing_scripts = [
                path for path in decision_scripts if not path.exists()
            ]
            if missing_scripts:
                raise FileNotFoundError(
                    "Configured decision scripts not found: "
                    + ", ".join(str(path) for path in missing_scripts)
                )

            backup_path = run_directory / (
                f"_ORIGINAL_dane_backup_{simulation_signature}.xlsx"
            )
            shutil.copy2(input_file, backup_path)

            try:
                shutil.copy2(all_output_path, input_file)
                print("Temporary dane.xlsx replaced by KNN-imputed ALL workbook.")

                for script_path in decision_scripts:
                    print(f"Running decision script: {script_path.name}")
                    completed_process = subprocess.run(
                        [sys.executable, str(script_path)],
                        cwd=project_root,
                        check=False,
                    )
                    if completed_process.returncode != 0:
                        raise RuntimeError(
                            f"Decision script '{script_path.name}' failed with "
                            f"exit code {completed_process.returncode}."
                        )
            finally:
                shutil.copy2(backup_path, input_file)
                backup_path.unlink(missing_ok=True)
                print("Original dane.xlsx restored.")

    print("\n========== KNN RUN FINISHED ==========")
    print(f"Run files: {run_directory}")
    print(f"Parameters history: {output_directory / 'KNN_parameters.xlsx'}")
    print(f"Imputation details: {details_path}")


if __name__ == "__main__":
    main()
